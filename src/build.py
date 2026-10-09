import json, html, os
from data import *
L="ABCDE"; E=html.escape
OUT=os.environ.get("OUT","out"); SITE=os.environ.get("SITE")=="1"; os.makedirs(OUT,exist_ok=True)
TITLE="成語習作 第1回"
STD,COL=(("xizuo-1-handout.docx","xizuo-1-handout-color.docx") if os.environ.get("SITE")=="1" else ("成語習作第1回_標準版.docx","成語習作第1回_彩色版.docx"))

# ---------- HTML ----------
def card(n,c):
    w,j,e,s,o,d=c
    rows=f'<p><b class="k">解釋</b>{E(j)}</p><p><b class="k">例句</b>{E(e)}</p>'
    if s: rows+=f'<p><b class="k s">相似</b>{E(s)}</p>'
    if o: rows+=f'<p><b class="k o">相反</b>{E(o)}</p>'
    if d: rows+=f'<p class="dn"><b class="k d">補充</b>{E(d)}</p>'
    return f'<div class="card"><h3><span>{n}</span>{E(w)}</h3>{rows}</div>'
cards="".join(card(i+1,c) for i,c in enumerate(I))
def mcq(n,q,opts,a):
    o="".join(f'<button type="button" class="opt" data-v="{L[i]}">({L[i]}) {E(t)}</button>' for i,t in enumerate(opts))
    return f'<div class="q" data-a="{a}" data-ans="{a}"><p>{n}. {E(q)}</p><div class="opts">{o}</div><div class="fb"></div></div>'

import random
rnd=random.Random(7); names=[c[0] for c in I]; GEN=[]
for i,c in enumerate(I):
    w,j,e,sm,op,x=c; pos=i%4; opts=None
    for kind,src,lab in (("o",op,"相反"),("s",sm,"相似")):
        if i%2==1 and src and opts is None:
            ans=src.split()[0]; bad=[n for n in names if n!=w and n not in sm.split() and n not in op.split()]
            opts=rnd.sample(bad,3); opts.insert(pos,ans); q=f"下列哪一個成語，與「{w}」的意思{'相反' if kind=='o' else '相近'}？"
    if opts is None:
        q=f"「{w}」的意思是："; opts=[k[1] for k in rnd.sample([k for k in I if k[0]!=w],3)]; opts.insert(pos,j)
    GEN.append((q,opts,L[pos]))

GENY=[]
for i,c in enumerate(I):
    w,j,e,sm,op,x=c
    if w not in e: e=e.replace('殃及無辜的人','殃及池魚，傷害到無辜的人')
    assert w in e,w
    bad=[n for n in names if n!=w and n not in sm.split() and n not in op.split()]
    opts=rnd.sample(bad,3); pos=(i+1)%4; opts.insert(pos,w)
    GENY.append((e.replace(w,"□□□□"),opts,L[pos]))
YY=YY+GENY
zx="".join(f'<div class="q" data-a="{a}" data-ans="{a}"><p>{i+1}. {E(q)}　<input type="text" maxlength="2" size="3" autocomplete="off"></p><div class="fb"></div></div>' for i,(q,a) in enumerate(ZX))
zh="".join(mcq(i+1,q,o,a) for i,(q,o,a) in enumerate(ZH))
yy="".join(mcq(i+1,q,o,a) for i,(q,o,a) in enumerate(YY))
ropts="".join(f'<option value="{L[i]}">({L[i]}) {E(t)}</option>' for i,t in enumerate(LR))
ll="".join(f'<div class="q" data-a="{LA[i]}" data-ans="{LA[i]}"><p>{i+1}. {E(t)} ⇄ <select><option value="">請選擇</option>{ropts}</select></p><div class="fb"></div></div>' for i,t in enumerate(LL))
wx="".join(f'<div class="q" data-a="{a}" data-ans="{a}"><p>{i+1}. {E(w)} → <select><option value="">請選擇</option><option>晴</option><option>陰</option><option>雨</option></select></p><div class="fb"></div></div>' for i,(w,a) in enumerate(WX))
gn="".join(mcq(i+1,q,o,a) for i,(q,o,a) in enumerate(GEN))
bt="".join(f'<p><b class="k">{E(a)}</b>{E(b)}</p>' for a,b in BT)

total=len(GEN)+len(ZX)+len(ZH)+len(LL)+len(YY)+len(WX)
ans=[]
ans.append("字形測驗："+"　".join(f"{i+1}.{a}" for i,(q,a) in enumerate(ZX)))
ans.append("綜合測驗："+"　".join(f"{i+1}.{a}" for i,(q,o,a) in enumerate(ZH)))
ans.append("連連看（意思相反）："+"　".join(f"{LL[i]}－{LR[ord(LA[i])-65]}" for i in range(5)))
ans.append("成語運用："+"　".join(f"{i+1}.{a}" for i,(q,o,a) in enumerate(YY)))
ans.append("成語氣象站："+"　".join(f"{w}（{a}）" for w,a in WX))
ans.append("綜合練習題："+"　".join(f"{i+1}.{a}" for i,(q,o,a) in enumerate(GEN)))
ansh="".join(f"<p>{E(x)}</p>" for x in ans)
CRUMB='<p><a href="../">← 成語典故</a>　<a href="static.html">講義版（可列印）</a></p>' if SITE else ''
page=f'''<!DOCTYPE html><html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{TITLE}｜互動版</title>
<style>
:root{{--bg:#fbf8f1;--fg:#2b2b2b;--card:#fff;--ac:#2f6f5e;--bd:#e2dccb;--ok:#d9f2e3;--bad:#fbdcdc}}
@media(prefers-color-scheme:dark){{:root{{--bg:#1d1f1e;--fg:#eee;--card:#262928;--ac:#7fd0b6;--bd:#3a3f3d;--ok:#1f4a37;--bad:#5a2b2b}}}}
body{{margin:0;background:var(--bg);color:var(--fg);font:16px/1.7 "Noto Sans TC","PingFang TC","Microsoft JhengHei",sans-serif}}
main{{max-width:960px;margin:auto;padding:16px}}h1{{color:var(--ac);margin:.3em 0}}h2{{border-left:6px solid var(--ac);padding-left:10px;margin-top:2em}}
nav a{{margin-right:12px;color:var(--ac)}}.bar{{position:sticky;top:0;background:var(--bg);padding:8px 0;border-bottom:1px solid var(--bd);z-index:5;display:flex;gap:10px;flex-wrap:wrap;align-items:center}}
button,select,input{{font:inherit}}.btn{{background:var(--ac);color:#fff;border:0;border-radius:8px;padding:6px 14px;cursor:pointer}}
.grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(290px,1fr));gap:12px}}.card{{background:var(--card);border:1px solid var(--bd);border-radius:12px;padding:10px 14px}}
.card h3{{margin:.2em 0;color:var(--ac)}}.card h3 span{{background:var(--ac);color:#fff;border-radius:50%;display:inline-block;width:1.7em;height:1.7em;text-align:center;font-size:.8em;line-height:1.7em;margin-right:8px}}
.card p{{margin:.3em 0}}.k{{background:var(--bd);border-radius:5px;padding:0 6px;margin-right:6px;font-size:.85em}}.s{{background:#cfe8ff;color:#123}}.o{{background:#ffe0c2;color:#321}}.d{{background:#e8dcff;color:#213}}.dn{{font-size:.92em}}
.q{{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:6px 12px;margin:8px 0}}.q p{{margin:.3em 0}}.opts{{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:6px}}
.opt{{text-align:left;background:transparent;color:var(--fg);border:1px solid var(--bd);border-radius:8px;padding:6px 10px;cursor:pointer}}.opt.sel{{border-color:var(--ac);background:var(--ok)}}
.q.ok{{background:var(--ok)}}.q.bad{{background:var(--bad)}}.fb{{font-size:.9em;font-weight:bold}}details{{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:8px 14px}}
.dl a{{display:inline-block;margin:6px 10px 6px 0;padding:8px 14px;border-radius:8px;background:var(--ac);color:#fff;text-decoration:none}}
</style></head><body><main>
{CRUMB}<h1>{TITLE}（中學生）</h1><p>40 個成語學習站、字形測驗、綜合測驗、連連看、成語運用（含 40 個例句應用題）、成語氣象站與 40 題綜合練習題。直接點選或輸入作答。</p>
<nav><a href="#learn">成語學習站</a><a href="#test">實力挑戰</a><a href="#extra">成語大補帖</a><a href="#answers">參考答案</a><a href="#dl">下載 Word</a></nav>
<p>班級：<input size="6"> 座號：<input size="4"> 姓名：<input size="8"></p>
<div class="bar"><label><input type="radio" name="m" value="live" checked> 即時對答案</label><label><input type="radio" name="m" value="sub"> 交卷模式</label>
<span id="sc">已作答 0／{total}　答對 0</span><button class="btn" id="go" type="button" hidden>交卷送出</button><button class="btn" id="rs" type="button">重新作答</button></div>
<h2 id="learn">成語學習站（共 40 個）</h2><div class="grid">{cards}</div>
<h2 id="test">實力挑戰</h2><h3>一、字形測驗（填入正確的字）</h3>{zx}
<h3>二、綜合測驗</h3>{zh}
<h3>三、連連看（選出意思相反的成語）</h3><p>右側選項：{"　".join(f"({L[i]}) {E(t)}" for i,t in enumerate(LR))}</p>{ll}
<h3>四、成語運用</h3>{yy}
<h3>五、成語氣象站（標示出心情是晴、陰或雨）</h3>{wx}
<h3>六、綜合練習題（共 40 題）</h3>{gn}
<h2 id="extra">成語大補帖：帥哥美女 V.S. 恐龍篇</h2><div class="card">{bt}</div>
<h2 id="answers">參考答案</h2><details><summary>顯示參考答案（老師用）</summary>{ansh}</details>
<h2 id="dl">下載 Word 檔</h2><div class="dl"><a href="{STD}">📄 標準版 Word 檔（黑白，適合列印）</a><a href="{COL}">🎨 彩色版 Word 檔（內容相同，依單元上色）</a></div>
</main><script>
const Q=[...document.querySelectorAll('.q')],T={total};
const val=q=>{{const i=q.querySelector('input[type=text]'),s=q.querySelector('select'),o=q.querySelector('.opt.sel');return i?i.value.trim():s?s.value:o?o.dataset.v:''}};
function mark(q){{const v=val(q);if(!v){{q.classList.remove('ok','bad');q.querySelector('.fb').textContent='';return}}const ok=v===q.dataset.a;q.classList.toggle('ok',ok);q.classList.toggle('bad',!ok);q.querySelector('.fb').textContent=ok?'✔ 答對了':'✘ 正確答案：'+q.dataset.ans}}
const mode=()=>document.querySelector('[name=m]:checked').value;
function stat(){{const d=Q.filter(q=>val(q)).length;const r=Q.filter(q=>q.classList.contains('ok')).length;document.getElementById('sc').textContent='已作答 '+d+'／'+T+'　答對 '+(mode()=='live'?r:'—')}}
function upd(q){{if(mode()=='live')mark(q);stat()}}
Q.forEach(q=>{{q.querySelectorAll('.opt').forEach(b=>b.onclick=()=>{{q.querySelectorAll('.opt').forEach(x=>x.classList.remove('sel'));b.classList.add('sel');upd(q)}});q.querySelectorAll('input,select').forEach(e=>e.addEventListener('input',()=>upd(q)))}});
document.querySelectorAll('[name=m]').forEach(r=>r.onchange=()=>{{document.getElementById('go').hidden=mode()=='live';Q.forEach(q=>{{q.classList.remove('ok','bad');q.querySelector('.fb').textContent=''}});if(mode()=='live')Q.forEach(mark);stat()}});
document.getElementById('go').onclick=()=>{{Q.forEach(q=>{{if(val(q))mark(q);else{{q.classList.add('bad');q.querySelector('.fb').textContent='未作答，正確答案：'+q.dataset.ans}}}});const r=Q.filter(q=>q.classList.contains('ok')).length;document.getElementById('sc').textContent='已作答 '+Q.filter(q=>val(q)).length+'／'+T+'　答對 '+r+'　（'+Math.round(r*100/T)+' 分）'}};
document.getElementById('rs').onclick=()=>{{Q.forEach(q=>{{q.classList.remove('ok','bad');q.querySelector('.fb').textContent='';q.querySelectorAll('.opt').forEach(x=>x.classList.remove('sel'));q.querySelectorAll('input,select').forEach(e=>e.value='')}});stat()}};
</script></body></html>'''
open(OUT+"/index.html","w",encoding="utf-8").write(page)


if SITE:
    def sq(n,q,o,a): return f'<div class="q"><p>{n}. {E(q)}</p><p class="o">'+"　".join(f"({L[k]}) {E(t)}" for k,t in enumerate(o))+'</p></div>'
    S=lambda lst:"".join(sq(i+1,q,o,a) for i,(q,o,a) in enumerate(lst))
    sc="".join(f'<div class="c"><b>{i+1}. {E(c[0])}</b><br>解釋：{E(c[1])}<br>例句：{E(c[2])}'+(f'<br>相似：{E(c[3])}' if c[3] else '')+(f'<br>相反：{E(c[4])}' if c[4] else '')+(f'<br>補充：{E(c[5])}' if c[5] else '')+'</div>' for i,c in enumerate(I))
    sz="".join(f"<p>{i+1}. {E(q)}：（　　）</p>" for i,(q,a) in enumerate(ZX))
    sl="<table>"+"".join(f"<tr><td>{LL[i]}</td><td>（　）</td><td>({L[i]}) {LR[i]}</td></tr>" for i in range(5))+"</table>"
    sw="".join(f"<p>{i+1}. {E(w)}（　　）</p>" for i,(w,a) in enumerate(WX))
    st=f'''<!DOCTYPE html><html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{TITLE}｜講義版</title>
<style>body{{font:15px/1.7 "Noto Sans TC","Microsoft JhengHei",sans-serif;color:#000;background:#fff;max-width:820px;margin:auto;padding:16px}}h1,h2{{color:#1f5c5a}}.c{{break-inside:avoid;margin:8px 0;padding:6px 0;border-bottom:1px solid #ccc}}.q{{break-inside:avoid;margin:6px 0}}.o{{margin:0 0 4px 1.5em}}table{{border-collapse:collapse}}td{{border:1px solid #888;padding:2px 12px}}.pb{{break-before:page}}@media print{{.np{{display:none}}}}</style></head><body>
<p class="np"><a href="../">← 成語典故</a>　<a href="./">互動版</a>　<a href="{STD}">Word 講義</a>　<a href="{COL}">Word 講義（彩色版）</a></p>
<h1>{TITLE}（中學生）</h1><p>班級：＿＿＿＿　座號：＿＿＿＿　姓名：＿＿＿＿＿＿　得分：＿＿＿＿</p>
<h2>壹、成語學習站</h2>{sc}<h2 class="pb">貳、實力挑戰</h2><h3>一、字形測驗</h3>{sz}<h3>二、綜合測驗</h3>{S(ZH)}<h3>三、連連看（意思相反）</h3>{sl}<h3>四、成語運用</h3>{S(YY)}<h3>五、成語氣象站（晴／陰／雨）</h3>{sw}<h3>六、綜合練習題</h3>{S(GEN)}
<h2>參、成語大補帖</h2>{bt}<h2 class="pb">參考答案</h2>{ansh}</body></html>'''
    open(OUT+"/static.html","w",encoding="utf-8").write(st)

# ---------- DOCX ----------
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
def mk(color,fn):
    d=Document(); sec=d.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    for m in ("left_margin","right_margin","top_margin","bottom_margin"): setattr(sec,m,Cm(2))
    st=d.styles["Normal"]; st.font.name="Microsoft JhengHei"; st.font.size=Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei")
    C=lambda rgb: RGBColor.from_string(rgb) if color else RGBColor(0,0,0)
    def P(parts,size=None,align=None,after=4):
        p=d.add_paragraph(); p.paragraph_format.space_after=Pt(after)
        if align: p.alignment=align
        for t,b,rgb in parts:
            r=p.add_run(t); r.bold=b; r.font.color.rgb=C(rgb) if rgb else RGBColor(0,0,0)
            r.font.name="Microsoft JhengHei"; r._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft JhengHei")
            if size: r.font.size=Pt(size)
        return p
    def H(t,lv=1,rgb="2F6F5E"):
        P([(t,True,rgb)],size=18 if lv==1 else 14,after=6)
    P([(TITLE+"（中學生）",True,"2F6F5E")],size=24,align=WD_ALIGN_PARAGRAPH.CENTER)
    P([("班級：＿＿＿＿　座號：＿＿＿＿　姓名：＿＿＿＿＿＿　得分：＿＿＿＿",False,None)])
    H("壹、成語學習站")
    for n,(w,j,e,s,o,x) in enumerate(I,1):
        P([(f"{n}. {w}",True,"B03A2E")],size=14,after=2)
        P([("解釋：",True,"1F5FA8"),(j,False,None)],after=1)
        P([("例句：",True,"1F5FA8"),(e,False,None)],after=1)
        if s: P([("相似：",True,"2E7D32"),(s,False,None)],after=1)
        if o: P([("相反：",True,"E65100"),(o,False,None)],after=1)
        if x: P([("補充：",True,"6A1B9A"),(x,False,None)],after=1)
        P([("",False,None)],after=2)
    d.add_page_break()
    H("貳、實力挑戰")
    H("一、字形測驗（填入正確的字）",2,"1F5FA8")
    for i,(q,a) in enumerate(ZX,1): P([(f"{i}. {q}：（　　　）",False,None)])
    H("二、綜合測驗",2,"1F5FA8")
    def M(lst):
        for i,(q,o,a) in enumerate(lst,1):
            P([(f"（　）{i}. {q}",False,None)],after=1)
            for k,t in enumerate(o): P([(f"　　({L[k]}) {t}",False,None)],after=0)
            P([("",False,None)],after=2)
    M(ZH)
    H("三、連連看（將意思相反的成語連起來）",2,"1F5FA8")
    t=d.add_table(rows=5,cols=3); t.style="Table Grid"
    for i in range(5):
        for j,v in enumerate([LL[i],"（　）",f"({L[i]}) {LR[i]}"]): t.cell(i,j).text=v
    P([("",False,None)])
    H("四、成語運用",2,"1F5FA8"); M(YY)
    H("五、成語氣象站（標示出心情：晴／陰／雨）",2,"1F5FA8")
    for i,(w,a) in enumerate(WX,1): P([(f"{i}. {w}（　　）",False,None)])
    H("六、綜合練習題（共 40 題）",2,"1F5FA8"); M(GEN)
    H("參、成語大補帖：帥哥美女 V.S. 恐龍篇")
    for a,b in BT: P([(a+"：",True,"6A1B9A"),(b,False,None)])
    d.add_page_break(); H("參考答案")
    for x in ans: P([(x,False,None)])
    d.save(OUT+"/"+fn)
mk(False,STD); mk(True,COL)
print("ok",total)
